"""Local parsers emit elements without making embedding or vision calls."""

import base64
import re
from pathlib import Path

from app.config import get_settings
from app.ingest.models import Element
from app.ingest.validation import InvalidFile
from app.log import log


def picture(
    blob: bytes, page: int, *, bbox=None, section="", mime="image/png"
) -> Element:
    return Element(
        type="image",
        content={"base64": base64.b64encode(blob).decode(), "mime": mime},
        page=page,
        bbox=bbox,
        section=section,
        source_method="image",
    )


def table(
    rows, page: int, *, section="", sheet=None, method="table_parser"
) -> Element | None:
    rows = [["" if cell is None else str(cell).strip() for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2:
        return None
    width = max(map(len, rows))
    rows = [row + [""] * (width - len(row)) for row in rows]
    return Element(
        type="table",
        content={"columns": rows[0], "rows": rows[1:], "title": section},
        page=page,
        sheet=sheet,
        section=section,
        source_method=method,
    )


def _pdf(path: Path, customer: bool):
    import pdfplumber
    import pymupdf

    cfg = get_settings()
    elements, scanned, processed = [], 0, 0
    with pymupdf.open(path) as doc:
        if doc.needs_pass:
            raise InvalidFile(cfg.knowledge["parsing"]["pdf_password_error"])
        total = len(doc)
        end = min(total, cfg.uploads["digital_pages_per_file"]) if customer else total
        with pdfplumber.open(path) as tables_doc:
            for number in range(end):
                page = doc[number]
                text = page.get_text("text").strip()
                digital = (
                    len(re.sub(r"\s", "", text))
                    >= cfg.knowledge["parsing"]["digital_min_chars"]
                )
                if digital:
                    elements.append(
                        Element(
                            type="text",
                            content=text,
                            page=number + 1,
                            source_method="text_layer",
                        )
                    )
                    for rows in tables_doc.pages[number].extract_tables():
                        item = table(rows, number + 1)
                        if item:
                            elements.append(item)
                    for info in page.get_image_info(xrefs=True):
                        xref = info.get("xref", 0)
                        if xref:
                            data = doc.extract_image(xref)
                            elements.append(
                                picture(
                                    data["image"],
                                    number + 1,
                                    bbox=list(info["bbox"]),
                                    mime=f"image/{data['ext']}",
                                )
                            )
                else:
                    if customer and scanned >= cfg.uploads["scanned_pages_per_file"]:
                        break
                    scanned += 1
                    pix = page.get_pixmap(dpi=cfg.ocr["render_dpi"])
                    elements.append(picture(pix.tobytes("png"), number + 1))
                processed += 1
                log.step("PARSE", "PDF page", page=number + 1, digital=digital)
    return elements, total, processed


def _pptx(path: Path, customer: bool):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    elements = []
    total = len(presentation.slides)
    end = (
        min(total, get_settings().uploads["digital_pages_per_file"])
        if customer
        else total
    )

    def shape_elements(shape, page):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                shape_elements(child, page)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            elements.append(
                picture(shape.image.blob, page, mime=shape.image.content_type)
            )
        elif shape.has_table:
            item = table(
                [[cell.text for cell in row.cells] for row in shape.table.rows], page
            )
            if item:
                elements.append(item)
        elif shape.has_text_frame and shape.text.strip():
            elements.append(
                Element(
                    type="text",
                    content=shape.text,
                    page=page,
                    source_method="text_layer",
                )
            )

    for number, slide in enumerate(presentation.slides):
        if number >= end:
            break
        for shape in slide.shapes:
            shape_elements(shape, number + 1)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                elements.append(
                    Element(
                        type="text",
                        content=notes,
                        page=number + 1,
                        section="Notes",
                        source_method="text_layer",
                    )
                )
    return elements, total, end


def _docx(path: Path, customer: bool):
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(path)
    elements, section = [], ""
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            section = paragraph.text
        if paragraph.text.strip():
            elements.append(
                Element(
                    type="text",
                    content=paragraph.text,
                    page=1,
                    section=section,
                    source_method="text_layer",
                )
            )
    for native in doc.tables:
        item = table([[cell.text for cell in row.cells] for row in native.rows], 1)
        if item:
            elements.append(item)
    for shape in doc.inline_shapes:
        for blip in shape._inline.iter(qn("a:blip")):
            part = doc.part.related_parts[blip.get(qn("r:embed"))]
            elements.append(picture(part.blob, 1, mime=part.content_type))
    return elements, 1, 1


def _grid(rows):
    # A title row is not a header; prefer the first populated, mostly textual row.
    scan = get_settings().knowledge["parsing"]["header_scan_rows"]
    rows = [list(row) for row in rows]
    for index, row in enumerate(rows[:scan]):
        values = [
            str(value).strip()
            for value in row
            if value is not None and str(value).strip()
        ]
        if (
            len(values) >= 2
            and sum(bool(re.search(r"[A-Za-z]", value)) for value in values)
            >= len(values) / 2
        ):
            return rows[index:]
    return rows


def _spreadsheet(path: Path, customer: bool):
    import openpyxl
    import pandas as pd

    elements = []
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(
            path,
            header=None,
            dtype=str,
            keep_default_na=False,
            sep=None,
            engine="python",
        )
        item = table(
            _grid(frame.values.tolist()),
            1,
            section=path.stem,
            sheet=path.stem,
            method="spreadsheet",
        )
        return ([item] if item else []), 1, 1
    book = openpyxl.load_workbook(path, data_only=True)
    end = (
        min(len(book.worksheets), get_settings().uploads["digital_pages_per_file"])
        if customer
        else len(book.worksheets)
    )
    try:
        for number, sheet in enumerate(book.worksheets[:end]):
            # Populate merged labels before detecting headers so their meaning is retained.
            for merged in list(sheet.merged_cells.ranges):
                value = sheet.cell(merged.min_row, merged.min_col).value
                sheet.unmerge_cells(str(merged))
                for row in sheet.iter_rows(
                    min_row=merged.min_row,
                    max_row=merged.max_row,
                    min_col=merged.min_col,
                    max_col=merged.max_col,
                ):
                    for cell in row:
                        cell.value = value
            item = table(
                _grid(sheet.values),
                number + 1,
                section=sheet.title,
                sheet=sheet.title,
                method="spreadsheet",
            )
            if item:
                elements.append(item)
        return elements, len(book.worksheets), end
    finally:
        book.close()


def parse_file(path: Path, *, customer: bool = False):
    ext = path.suffix.lower()
    with log.span("PARSE", "parse file", extension=ext):
        try:
            if ext == ".pdf":
                return _pdf(path, customer)
            if ext == ".pptx":
                return _pptx(path, customer)
            if ext == ".docx":
                return _docx(path, customer)
            if ext in {".xlsx", ".csv"}:
                return _spreadsheet(path, customer)
            if ext in {".md", ".txt"}:
                elements, section, buffer = [], "", []
                for line in path.read_text(encoding="utf-8-sig").splitlines():
                    if line.startswith("#"):
                        if buffer:
                            elements.append(
                                Element(
                                    type="text",
                                    content="\n".join(buffer),
                                    page=1,
                                    section=section,
                                    source_method="text_layer",
                                )
                            )
                        section, buffer = line.lstrip("# "), []
                    else:
                        buffer.append(line)
                if buffer:
                    elements.append(
                        Element(
                            type="text",
                            content="\n".join(buffer),
                            page=1,
                            section=section,
                            source_method="text_layer",
                        )
                    )
                return elements, 1, 1
            return (
                [
                    picture(
                        path.read_bytes(),
                        1,
                        mime=f"image/{'jpeg' if ext in {'.jpg', '.jpeg'} else ext[1:]}",
                    )
                ],
                1,
                1,
            )
        except InvalidFile:
            raise
        except Exception as exc:
            log.step("PARSE", "failed", error=type(exc).__name__)
            raise InvalidFile(
                get_settings().knowledge["parsing"]["corrupt_error"]
            ) from None
