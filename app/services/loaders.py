"""Turn an uploaded file (bytes) into Sections: text blocks with a human-readable location.

Supported: .pdf .docx .pptx .txt .md .csv .tsv .xlsx

Design rules
- Every Section carries a `location` ("p.3", "Slide 4 – Pricing", "Sheet 'Sales' · rows 21–40")
  that the chat UI and citations show to the user, and a `heading` (breadcrumb such as
  "Q3 Review > Slide 4: Pricing") that is added to the text *only when embedding*, so the
  vector carries the context without cluttering the displayed text.
- Tables are never flattened into loose text. Each table row becomes "Column: value; ..."
  and the column names are repeated in every chunk, so the LLM always knows which header a
  number belongs to. Chunks hold whole rows only.
- Spreadsheet-style tables also get one "summary" section (row count, column types, sums,
  min/max, value counts) computed over ALL rows, because similarity search alone cannot
  answer "what is the total" or "how many are failed".
- Anything we could not read (images, row caps, hidden sheets...) becomes a warning string
  that is stored on the document and shown in the UI.
"""
from __future__ import annotations

import codecs
import csv
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import PurePath

from app.config import get_settings

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv", ".tsv", ".xlsx"}

MAX_READ_ROWS = 200_000      # stop reading a table after this many rows (statistics cover what was read)
MAX_COLUMNS = 200
MAX_PDF_PAGES = 400
SLIDE_ATOMIC_CHARS = 3000    # a slide shorter than this stays in one chunk
MAX_CELL_CHARS = 1000
MAX_STAT_COLUMNS = 40
MAX_WARNINGS = 12


class UnsupportedFileType(ValueError):
    pass


@dataclass
class Section:
    text: str                 # what the LLM and the user see
    location: str             # label used in citations
    page: int | None = None   # kept for PDFs (backwards compatible)
    kind: str = "text"        # text | table | table_summary | slide
    heading: str = ""         # embedding-only context prefix
    atomic: bool = False      # already sized; the chunker must not re-split it


@dataclass
class LoadResult:
    sections: list[Section] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def warn(self, msg: str) -> None:
        if msg not in self.warnings and len(self.warnings) < MAX_WARNINGS:
            self.warnings.append(msg)


# ---------------------------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------------------------
def clean_text(text: str) -> str:
    text = text.replace("\x00", "")           # null bytes break Mongo/LLM calls
    text = re.sub(r"[ \t]+", " ", text)       # collapse runs of spaces/tabs
    text = re.sub(r"\n{3,}", "\n\n", text)    # at most one blank line in a row
    return text.strip()


def _need(module: str, package: str):
    try:
        return __import__(module, fromlist=["_"])
    except ImportError as e:
        raise RuntimeError(f"Reading this file type needs the '{package}' package. Run: uv add {package}") from e


def _decode(data: bytes) -> str:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _cell_text(v) -> str:
    """Any spreadsheet/table cell value -> a clean one-line string."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0, 0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat(timespec="minutes")
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        return format(v, ".10g")
    s = re.sub(r"\s+", " ", str(v).replace("\x00", "")).strip()
    return s[:MAX_CELL_CHARS] + "…" if len(s) > MAX_CELL_CHARS else s


def _breadcrumb(path: list[tuple[int, str]]) -> str:
    return " > ".join(t for _, t in path)


def _push_heading(path: list[tuple[int, str]], level: int, text: str) -> list[tuple[int, str]]:
    return [p for p in path if p[0] < level] + [(level, text)]


# ---------------------------------------------------------------------------------------------
# tables (shared by every format)
# ---------------------------------------------------------------------------------------------
def _unique_headers(raw: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, h in enumerate(raw, start=1):
        h = h or f"Column {i}"
        n = seen.get(h, 0)
        seen[h] = n + 1
        out.append(h if n == 0 else f"{h} ({n + 1})")
    return out


def _prepare_grid(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """Drop empty rows/columns, make the grid rectangular, split header from body."""
    rows = [[str(c) for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return [], []
    width = min(max(len(r) for r in rows), MAX_COLUMNS)
    rows = [(r + [""] * width)[:width] for r in rows]
    keep = [j for j in range(width) if any(r[j] for r in rows)]
    rows = [[r[j] for j in keep] for r in rows]
    return _unique_headers(rows[0]), rows[1:]


def _row_line(headers: list[str], row: list[str], n: int) -> str | None:
    parts = [f"{h}: {v}" for h, v in zip(headers, row) if v]
    return f"Row {n}: " + "; ".join(parts) if parts else None


def _table_sections(headers, body, *, name, base, heading, total=None, page=None, kind="table") -> list[Section]:
    """Group rows into chunks of whole rows; every chunk repeats the table name and columns."""
    budget = get_settings().table_chunk_chars
    prefix = f"Table: {name}\nColumns: {' | '.join(headers)}\n"
    avail = max(300, budget - len(prefix) - 40)
    groups: list[list[tuple[int, str]]] = []
    cur: list[tuple[int, str]] = []
    size = 0
    for n, row in enumerate(body, start=1):
        line = _row_line(headers, row, n)
        if line is None:
            continue
        if cur and size + len(line) + 1 > avail:
            groups.append(cur)
            cur, size = [], 0
        cur.append((n, line))
        size += len(line) + 1
    if cur:
        groups.append(cur)

    total = total or len(body)
    out = []
    for g in groups:
        a, b = g[0][0], g[-1][0]
        span = f"Rows {a}–{b}" if a != b else f"Row {a}"
        text = f"{prefix}{span} of {total}:\n" + "\n".join(line for _, line in g)
        loc = " · ".join(x for x in (base, span.lower()) if x)
        out.append(Section(text=text, location=loc, page=page, kind=kind, heading=heading, atomic=True))
    return out


_NUM = re.compile(r"^[-+]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}")
_ID_NAME = re.compile(r"(?i)(^|[\s_\-])(id|idx|index|uuid|zip|pin|phone|mobile)($|[\s_\-])")


def _fmt_num(x: float) -> str:
    s = f"{x:.2f}".rstrip("0").rstrip(".")
    return s or "0"


def _as_numbers(values: list[str]) -> list[float] | None:
    nums = []
    for v in values:
        t = re.sub(r"^[\$€£₹]\s*", "", v)
        if _NUM.match(t):
            nums.append(float(t.replace(",", "")))
    if values and len(nums) / len(values) >= 0.9 and not any(re.match(r"^0\d", v) for v in values):
        return nums
    return None


def _column_stats(name: str, values: list[str], n_rows: int) -> str | None:
    if not values:
        return None
    blanks = n_rows - len(values)
    tail = f", {blanks} blank" if blanks > 0 else ""

    if not _ID_NAME.search(name):
        nums = _as_numbers(values)
        if nums:
            return (f"- {name} (number): {len(nums)} values, sum {_fmt_num(sum(nums))}, min {_fmt_num(min(nums))}, "
                    f"max {_fmt_num(max(nums))}, mean {_fmt_num(sum(nums) / len(nums))}{tail}")

    dates = [v[:10] for v in values if _ISO_DATE.match(v)]
    if dates and len(dates) / len(values) >= 0.9:
        return f"- {name} (date): earliest {min(dates)}, latest {max(dates)}{tail}"

    counts = Counter(values)
    k = len(counts)
    if k <= 12:
        top = ", ".join(f"{v[:40]} ({c})" for v, c in counts.most_common())
        return f"- {name} (text): {k} distinct value{'s' if k != 1 else ''}: {top}{tail}"
    return f"- {name} (text): {k} distinct values{tail}"


def _summary_section(headers, body, *, name, base, heading, indexed: int) -> Section:
    n = len(body)
    lines = [
        f"Table summary: {name}",
        f"Total data rows: {n} (row numbers count the rows below the header row)",
    ]
    if indexed < n:
        lines.append(f"Only the first {indexed} rows are searchable; the statistics below cover all {n} rows.")
    lines.append(f"Columns ({len(headers)}): " + ", ".join(headers[: MAX_STAT_COLUMNS * 2]))
    stats = []
    for j, h in enumerate(headers[:MAX_STAT_COLUMNS]):
        s = _column_stats(h, [r[j] for r in body if r[j]], n)
        if s:
            stats.append(s)
    if stats:
        lines.append("Column statistics (computed over all rows):")
        lines.extend(stats)
    sample = [ln for ln in (_row_line(headers, r, i) for i, r in enumerate(body[:3], start=1)) if ln]
    if sample:
        lines.append("Sample rows:")
        lines.extend(sample)
    loc = " · ".join(x for x in (base, "summary") if x)
    return Section(text="\n".join(lines)[:7000], location=loc, kind="table_summary", heading=heading, atomic=True)


def _split_preamble(rows: list[list[str]]) -> tuple[list[list[str]], list[list[str]]]:
    """Title/notes rows above the real header (common in spreadsheets) are split off."""
    counts = [(i, sum(1 for c in r if c)) for i, r in enumerate(rows) if any(r)]
    if not counts:
        return [], rows
    window = counts[:10]
    need = max(1, -(-max(c for _, c in window) * 3 // 5))  # ceil(60% of the widest row)
    for i, c in window:
        if c >= need:
            return rows[:i], rows[i:]
    return [], rows


def _grid_sections(rows, *, name, base, heading, result: LoadResult, summary: bool = True) -> list[Section]:
    """Full treatment for a spreadsheet-style table: notes + summary + row-group chunks."""
    sections: list[Section] = []
    preamble, rows = _split_preamble(rows)
    notes = " ".join(c for r in preamble for c in r if c)
    if notes:
        sections.append(Section(text=f"{name} (notes above the table): {notes}",
                                location=" · ".join(x for x in (base, "notes") if x), heading=heading))
    headers, body = _prepare_grid(rows)
    if not body:
        result.warn(f"'{name}' has a header but no data rows.")
        return sections
    cap = get_settings().max_table_rows
    indexed = body[:cap]
    if len(body) > cap:
        result.warn(f"'{name}' has {len(body)} rows; only the first {cap} are searchable "
                    f"(the summary statistics cover all {len(body)}).")
    if summary:
        sections.append(_summary_section(headers, body, name=name, base=base, heading=heading, indexed=len(indexed)))
    sections += _table_sections(headers, indexed, name=name, base=base, heading=heading, total=len(body))
    return sections


def _small_table_sections(rows, *, name, base, heading, page=None, result: LoadResult) -> list[Section]:
    """Tables embedded in documents/slides: no summary, just header-preserving row chunks."""
    headers, body = _prepare_grid(rows)
    if not body:
        # a one-row "table" is layout, not data: keep it as plain text
        text = " | ".join(c for r in rows for c in r if c)
        return [Section(text=text, location=base, page=page, heading=heading)] if text else []
    return _table_sections(headers, body, name=name, base=base, heading=heading, page=page)


# ---------------------------------------------------------------------------------------------
# format loaders: each appends Sections to `result`
# ---------------------------------------------------------------------------------------------
def _load_text(data: bytes, title: str, result: LoadResult) -> None:
    text = clean_text(_decode(data))
    if text:
        result.sections.append(Section(text=text, location="Text", heading=title))


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _md_cells(line: str) -> list[str]:
    return [_cell_text(c) for c in line.strip().strip("|").split("|")]


def _load_markdown(data: bytes, title: str, result: LoadResult) -> None:
    lines = _decode(data).replace("\x00", "").replace("\r\n", "\n").split("\n")
    path: list[tuple[int, str]] = []
    buf: list[str] = []
    tables = 0
    in_fence = False

    def flush():
        text = "\n".join(buf).strip()
        buf.clear()
        if text:
            crumb = _breadcrumb(path)
            result.sections.append(Section(text=text, location=crumb or "Document",
                                           heading=f"{title} > {crumb}" if crumb else title))

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            in_fence = not in_fence
            buf.append(line)
        elif in_fence:
            buf.append(line)
        elif (m := _MD_HEADING.match(line)):
            flush()
            path = _push_heading(path, len(m.group(1)), m.group(2))
            buf.append(m.group(2))
        elif line.lstrip().startswith("|") and i + 1 < len(lines) and _MD_SEP.match(lines[i + 1]):
            flush()
            rows = [_md_cells(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_md_cells(lines[i]))
                i += 1
            tables += 1
            crumb = _breadcrumb(path)
            base = f"{crumb} · table {tables}" if crumb else f"Table {tables}"
            result.sections += _small_table_sections(
                rows, name=f"Table {tables}", base=base, heading=f"{title} > {base}", result=result)
            continue
        else:
            buf.append(line)
        i += 1
    flush()


def _load_delimited(data: bytes, title: str, result: LoadResult, ext: str) -> None:
    text = _decode(data)
    if ext == ".tsv":
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    csv.field_size_limit(10_000_000)
    rows: list[list[str]] = []
    try:
        for i, r in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delim)):
            if i >= MAX_READ_ROWS:
                result.warn(f"'{title}' is very large; only the first {MAX_READ_ROWS:,} rows were read.")
                break
            rows.append([_cell_text(c) for c in r[:MAX_COLUMNS]])
    except csv.Error as e:
        result.warn(f"'{title}': stopped reading at a malformed line ({e}).")
    result.sections += _grid_sections(rows, name=title, base="", heading=title, result=result)


def _load_xlsx(data: bytes, title: str, result: LoadResult) -> None:
    openpyxl = _need("openpyxl", "openpyxl")
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                result.warn(f"Hidden sheet '{ws.title}' was skipped.")
                continue
            try:
                ws.reset_dimensions()  # some generators write a wrong sheet size
            except Exception:  # noqa: BLE001
                pass
            rows: list[list[str]] = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_READ_ROWS:
                    result.warn(f"Sheet '{ws.title}' is very large; only the first {MAX_READ_ROWS:,} rows were read.")
                    break
                rows.append([_cell_text(c) for c in r[:MAX_COLUMNS]])
            if not any(any(r) for r in rows):
                result.warn(f"Sheet '{ws.title}' is empty.")
                continue
            base = f"Sheet '{ws.title}'"
            result.sections += _grid_sections(rows, name=ws.title, base=base,
                                              heading=f"{title} > {base}", result=result)
    finally:
        wb.close()


def _load_docx(data: bytes, title: str, result: LoadResult) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    path: list[tuple[int, str]] = []
    buf: list[str] = []
    tables = 0

    def flush():
        text = "\n".join(buf).strip()
        buf.clear()
        if text:
            crumb = _breadcrumb(path)
            result.sections.append(Section(text=text, location=crumb or "Document",
                                           heading=f"{title} > {crumb}" if crumb else title))

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style == "Title" or style.startswith("Heading"):
                m = re.search(r"(\d+)", style)
                flush()
                path = _push_heading(path, int(m.group(1)) if m else 0, text[:120])
                buf.append(text)
            elif style.startswith("List"):
                buf.append("- " + text)
            else:
                buf.append(text)
        elif child.tag == qn("w:tbl"):
            flush()
            tables += 1
            t = Table(child, doc)
            rows = []
            for r in t.rows:
                cells, last = [], None
                for c in r.cells:
                    if c._tc is last:  # horizontally merged cells repeat the same cell
                        continue
                    last = c._tc
                    cells.append(_cell_text(c.text))
                rows.append(cells)
            crumb = _breadcrumb(path)
            base = f"{crumb} · table {tables}" if crumb else f"Table {tables}"
            result.sections += _small_table_sections(
                rows, name=f"Table {tables}", base=base, heading=f"{title} > {base}", result=result)
    flush()


def _load_pdf_plain(data: bytes, title: str, result: LoadResult) -> None:
    from pypdf import PdfReader

    for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
        text = clean_text(page.extract_text() or "")
        if text:
            result.sections.append(Section(text=text, location=f"p.{i}", page=i, heading=f"{title} > p.{i}"))


def _load_pdf(data: bytes, title: str, result: LoadResult) -> None:
    try:
        import pdfplumber
        pdf = pdfplumber.open(io.BytesIO(data))
    except ImportError:
        result.warn("Tables in PDFs are read as plain text (install 'pdfplumber' for better table extraction).")
        return _load_pdf_plain(data, title, result)
    except Exception:  # noqa: BLE001 - unreadable by pdfplumber: try the simpler reader
        return _load_pdf_plain(data, title, result)

    with pdf:
        pages = pdf.pages
        if len(pages) > MAX_PDF_PAGES:
            result.warn(f"Only the first {MAX_PDF_PAGES} of {len(pages)} pages were read.")
            pages = pages[:MAX_PDF_PAGES]
        for i, page in enumerate(pages, start=1):
            tables: list[list[list[str]]] = []
            try:
                found = []
                for t in page.find_tables():
                    grid = [[_cell_text(c) for c in r] for r in t.extract()]
                    if len(grid) >= 2 and max(len(r) for r in grid) >= 2:
                        found.append((t, grid))
                work = page
                for t, _ in found:
                    work = work.outside_bbox(t.bbox)  # text outside tables; tables are handled separately
                text = clean_text(work.extract_text() or "")
                tables = [g for _, g in found]
            except Exception:  # noqa: BLE001 - fall back to plain page text
                text = clean_text(page.extract_text() or "")
            if text:
                result.sections.append(Section(text=text, location=f"p.{i}", page=i, heading=f"{title} > p.{i}"))
            for k, grid in enumerate(tables, start=1):
                base = f"p.{i} · table {k}"
                result.sections += _small_table_sections(
                    grid, name=f"Table {k} on page {i}", base=base, heading=f"{title} > {base}", page=i, result=result)


def _load_pptx(data: bytes, title: str, result: LoadResult) -> None:
    _need("pptx", "python-pptx")
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    skip_ph = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}

    def walk(shapes):
        for sh in sorted(shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh

    def alt_text(sh) -> str:
        try:
            alt = (sh._element.xpath(".//p:cNvPr/@descr") or [""])[0].strip()
        except Exception:  # noqa: BLE001
            return ""
        # tools auto-fill alt text with the file name ("image.png", "IMG_1234.jpg", "Picture 3"): noise
        if re.fullmatch(r"(?i)\S+\.(png|jpe?g|gif|bmp|svg|tiff?|webp|emf|wmf)", alt) \
                or re.fullmatch(r"(?i)(picture|image|graphic|photo)\s*\d*", alt):
            return ""
        return alt

    def chart_lines(sh) -> list[str]:
        out = []
        try:
            chart = sh.chart
            name = ""
            if chart.has_title and chart.chart_title.has_text_frame:
                name = chart.chart_title.text_frame.text.strip()
            for plot in chart.plots:
                cats = [str(c) for c in plot.categories]
                for s in plot.series:
                    vals = list(s.values)
                    pairs = ", ".join(f"{c}={_cell_text(v)}" for c, v in zip(cats, vals)) if cats else \
                        ", ".join(_cell_text(v) for v in vals)
                    out.append(f"Chart{f' ({name})' if name else ''} – {s.name}: {pairs}")
        except Exception:  # noqa: BLE001
            out.append("Chart (its data could not be read)")
        return out

    prs = Presentation(io.BytesIO(data))
    saw_images = False
    unreadable: list[int] = []

    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        slide_title = title_shape.text_frame.text.strip().replace("\x0b", " ") if title_shape is not None else ""
        body: list[str] = []
        n_tables = 0
        images = 0

        for sh in walk(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if sh.is_placeholder and sh.placeholder_format.type in skip_ph:
                continue
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = p.text.replace("\x0b", " ").strip()
                    if t:
                        body.append(("  " * min(p.level, 3) + "- " if p.level > 0 else "") + t)
            elif sh.has_table:
                n_tables += 1
                rows = [[_cell_text(c.text) for c in r.cells] for r in sh.table.rows]
                headers, trows = _prepare_grid(rows)
                if trows:
                    body.append(f"Table {n_tables}: columns {' | '.join(headers)}")
                    body += [ln for ln in (_row_line(headers, r, k) for k, r in enumerate(trows, start=1)) if ln]
                else:
                    body.append(" | ".join(c for r in rows for c in r if c))
            elif sh.has_chart:
                body += chart_lines(sh)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images += 1
                alt = alt_text(sh)
                if alt:
                    body.append(f"Image: {alt}")

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            body.append(f"Speaker notes: {notes}")
        if images:
            saw_images = True

        if not slide_title and not body:
            if images:
                unreadable.append(i)
            continue

        label = f"Slide {i}" + (f": {slide_title}" if slide_title else "")
        text = label + ("\n" + "\n".join(body) if body else "")
        loc = f"Slide {i}" + (f" – {slide_title[:50]}" if slide_title else "")
        result.sections.append(Section(text=text, location=loc, kind="slide", heading=f"{title} > {label}",
                                       atomic=len(text) <= SLIDE_ATOMIC_CHARS))

    if saw_images:
        result.warn("Images in slides are not read (no OCR); only their alt text is indexed.")
    if unreadable:
        result.warn("Slides with no readable text: " + ", ".join(map(str, unreadable[:10])))


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".pptx": _load_pptx,
    ".txt": _load_text,
    ".md": _load_markdown,
    ".xlsx": _load_xlsx,
    ".csv": lambda d, t, r: _load_delimited(d, t, r, ".csv"),
    ".tsv": lambda d, t, r: _load_delimited(d, t, r, ".tsv"),
}


def load_file(filename: str, data: bytes) -> LoadResult:
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    result = LoadResult()
    try:
        _LOADERS[ext](data, PurePath(filename).stem, result)
    except RuntimeError:
        raise  # e.g. a missing optional package: the message already says what to install
    except Exception as e:  # noqa: BLE001 - corrupt, encrypted or malformed file
        raise ValueError(
            f"Could not read this {ext} file. It may be corrupt, password-protected or not a real "
            f"{ext} file ({type(e).__name__})."
        ) from e
    if not result.sections:
        raise ValueError(
            "No extractable text found (a scanned PDF or an image-only deck would need OCR)."
            + (" " + " ".join(result.warnings) if result.warnings else "")
        )
    return result