"""Phase 3 isolation, attachment lifecycle and ingestion limits without live services."""
import asyncio
import copy
import io
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

from bson import ObjectId
from pymongo.errors import DuplicateKeyError
import pytest

from test_phase2 import Collection, Cursor, world
from app.config import get_settings
from app.routers import chat, attachments as attachment_routes
from app.services import attachments, document_store, ingestion, memory, retrieval
from app.services.loaders import load_file
from app.services.router import Plan


class Mongo(Collection):
    async def insert_one(self, row):
        if row.get("scope") == "user":
            key = {field: row[field] for field in ("scope", "owner_id", "session_id", "file_hash")}
            if any(self.matches(existing, key) for existing in self.rows):
                raise DuplicateKeyError("Duplicate scoped key")
        self.rows.append(copy.deepcopy(row))

    async def insert_many(self, rows):
        self.rows.extend(copy.deepcopy(rows))

    async def aggregate(self, pipeline):
        self.pipeline = copy.deepcopy(pipeline)
        stage = pipeline[0]["$vectorSearch"]
        rows = [copy.deepcopy(row) for row in self.rows if self.matches(row, stage["filter"])]
        for row in rows:
            row["score"] = row.get("score", .9)
        rows.sort(key=lambda row: row["score"], reverse=True)
        return Cursor(rows[:stage["limit"]])


@pytest.fixture
def bank_files(world, monkeypatch):
    docs, chunks = Mongo(), Mongo()
    world.docs, world.chunks = docs, chunks
    for module in (retrieval, document_store, ingestion, memory, attachments, attachment_routes):
        if hasattr(module, "documents_col"):
            monkeypatch.setattr(module, "documents_col", lambda: docs)
        if hasattr(module, "chunks_col"):
            monkeypatch.setattr(module, "chunks_col", lambda: chunks)
    monkeypatch.setattr(document_store, "sessions_col", lambda: world.sessions)
    world.answer_sources, world.embedding_batches = [], []
    async def embed_documents(texts):
        world.embedding_batches.append(texts)
        return [[1., 0.] for _ in texts]
    async def embed_query(text): return [1., 0.]
    async def plan(history, message, names): return Plan(question=message)
    async def answer(history, question, sources):
        world.answer_sources.append(copy.deepcopy(sources))
        assert sources
        yield "Document fact [1]."
    monkeypatch.setattr(ingestion, "embed_documents", embed_documents)
    monkeypatch.setattr(retrieval, "embed_query", embed_query)
    monkeypatch.setattr(chat, "plan_turn", plan)
    monkeypatch.setattr(chat, "ready_document_names", retrieval.ready_document_names)
    monkeypatch.setattr(chat, "stream_answer", answer)
    return world


def new_session(state, user="a"):
    return state.client.post("/sessions", headers=state.headers(user)).json()["session_id"]


def upload(state, sid, data=b"Private document fact.", name="private.txt", user="a"):
    response = state.client.post(f"/sessions/{sid}/attachments", headers=state.headers(user),
        files=[("files", (name, data, "application/octet-stream"))])
    assert response.status_code == 202, response.text
    return response.json()[0]


def test_upload_propagates_scope_and_keeps_objectids(bank_files):
    state = bank_files
    sid = new_session(state)
    result = upload(state, sid)
    assert result["status"] == "processing"
    doc = state.docs.rows[0]
    assert isinstance(doc["_id"], ObjectId)
    assert {key: doc[key] for key in ("scope", "owner_id", "session_id")} == document_store.user_scope("a", sid)
    assert doc["status"] == "ready"
    assert state.chunks.rows
    for row in state.chunks.rows:
        assert row["scope"] == "user" and row["owner_id"] == "a" and row["session_id"] == sid
        assert row["doc_id"] == result["doc_id"] and row["expires_at"] == doc["expires_at"]
    listed = state.client.get(f"/sessions/{sid}/attachments", headers=state.headers("a")).json()
    assert listed[0]["id"] == result["doc_id"] and listed[0]["status"] == "ready"
    assert "file_hash" not in listed[0] and "owner_id" not in listed[0]


def test_deduplication_includes_owner_and_session(bank_files):
    state = bank_files
    s1, s2, sb = new_session(state), new_session(state), new_session(state, "b")
    first = upload(state, s1)
    duplicate = upload(state, s1)
    second_chat = upload(state, s2)
    other_user = upload(state, sb, user="b")
    assert duplicate["duplicate"] is True and duplicate["doc_id"] == first["doc_id"]
    assert len({first["doc_id"], second_chat["doc_id"], other_user["doc_id"]}) == 3
    assert state.sessions.rows[0]["attachment_count"] == 1


@pytest.mark.parametrize("stream", [False, True])
def test_foreign_or_wrong_chat_attachment_is_404_before_stream(bank_files, stream):
    state = bank_files
    sa, sa2, sb = new_session(state), new_session(state), new_session(state, "b")
    row = upload(state, sa)
    for sid, user in ((sa2, "a"), (sb, "b")):
        response = state.client.post("/chat", headers=state.headers(user), json={
            "message":"Read this file", "session_id":sid, "attachment_ids":[row["doc_id"]], "stream":stream})
        assert response.status_code == 404
    assert state.client.get(f"/sessions/{sa}/attachments", headers=state.headers("b")).status_code == 404
    assert state.client.post(f"/sessions/{sa}/attachments", headers=state.headers("b"),
        files=[("files", ("a.txt", b"hello"))]).status_code == 404
    assert state.answer_sources == []


def test_retrieval_filter_and_followups_never_include_other_files(bank_files):
    state = bank_files
    sa, sa2, sb = new_session(state), new_session(state), new_session(state, "b")
    own = upload(state, sa, b"Own fact.", "own.txt")
    upload(state, sa2, b"Other chat fact.", "other-chat.txt")
    upload(state, sb, b"Other user fact.", "other-user.txt", user="b")
    shared, _ = asyncio.run(document_store.register_document("shared.txt", "text/plain", b"Shared fact.", scope="shared"))
    asyncio.run(document_store.ingest_scoped(str(shared["_id"]), "shared.txt", b"Shared fact.", scope="shared"))
    # Legacy unscoped records cannot become public just because an index was rebuilt.
    legacy_id = ObjectId()
    state.docs.rows.append({"_id":legacy_id, "filename":"legacy.txt", "status":"ready", "created_at":datetime.now(timezone.utc)})
    state.chunks.rows.append({"doc_id":str(legacy_id), "text":"legacy secret", "score":1.})
    names = asyncio.run(retrieval.ready_document_names(owner_id="a", session_id=sa))
    assert set(names) == {"own.txt", "shared.txt"}
    for ids in ([own["doc_id"]], []):
        response = state.client.post("/chat", headers=state.headers("a"), json={
            "message":"What does the file say?", "session_id":sa, "attachment_ids":ids, "stream":False})
        assert response.status_code == 200
        assert {row["filename"] for row in response.json()["sources"]} == {"own.txt", "shared.txt"}
    filter_ = state.chunks.pipeline[0]["$vectorSearch"]["filter"]
    assert filter_["$and"][0] == {"$or":[{"scope":"shared"}, {"scope":"user", "owner_id":"a", "session_id":sa}]}
    assert set(filter_["$and"][1]["doc_id"]["$in"]) == {own["doc_id"], str(shared["_id"])}
    assert state.messages.rows[0]["attachments"] == [{"id":own["doc_id"], "filename":"own.txt"}]


def test_deleted_chat_removes_only_its_files_and_chunks(bank_files):
    state = bank_files
    sa, sb = new_session(state), new_session(state, "b")
    own, other = upload(state, sa), upload(state, sb, user="b")
    response = state.client.delete(f"/sessions/{sa}", headers=state.headers("a"))
    assert response.status_code == 200
    assert [str(row["_id"]) for row in state.docs.rows] == [other["doc_id"]]
    assert all(row["doc_id"] != own["doc_id"] for row in state.chunks.rows)


def test_expiry_refreshed_with_session_activity(bank_files):
    state = bank_files
    sid = new_session(state)
    upload(state, sid)
    old = datetime.now(timezone.utc) - timedelta(days=10)
    state.sessions.rows[0]["updated_at"] = old
    for row in state.docs.rows + state.chunks.rows:
        row["expires_at"] = old + timedelta(days=get_settings().retention_days)
    asyncio.run(memory.save_message(sid, "user", "a new message", user_id="a"))
    expiry = state.sessions.rows[0]["updated_at"] + timedelta(days=get_settings().retention_days)
    assert all(row["expires_at"] == expiry for row in state.docs.rows + state.chunks.rows)


def test_upload_count_size_type_and_empty_limits(bank_files, monkeypatch):
    state = bank_files
    sid = new_session(state)
    too_many = state.client.post(f"/sessions/{sid}/attachments", headers=state.headers("a"),
        files=[("files", (f"{i}.txt", b"x")) for i in range(4)])
    assert too_many.status_code == 400 and state.docs.rows == []
    for name, data in (("bad.exe", b"x"), ("empty.txt", b"")):
        assert upload(state, sid, data, name)["status"] == "rejected"
    monkeypatch.setenv("USER_UPLOAD_MAX_MB", "1")
    get_settings.cache_clear()
    assert upload(state, sid, b"x" * (1024*1024+1))["status"] == "rejected"
    for i in range(5):
        assert upload(state, sid, f"file {i}".encode(), f"{i}.txt")["status"] == "processing"
    assert upload(state, sid, b"sixth")["status"] == "rejected"
    assert upload(state, sid, b"file 0", "0.txt")["duplicate"] is True
    assert len(state.docs.rows) == 5


def test_concurrent_uploads_cannot_exceed_five_slots(bank_files):
    state = bank_files
    sid = new_session(state)
    async def go():
        return await asyncio.gather(*(document_store.register_document(f"{i}.txt", "text/plain", str(i).encode(),
            scope="user", owner_id="a", session_id=sid) for i in range(10)), return_exceptions=True)
    results = asyncio.run(go())
    assert sum(isinstance(row, tuple) for row in results) == 5
    assert state.sessions.rows[0]["attachment_count"] == 5
    assert len(state.docs.rows) == 5


def test_chunk_and_table_limits_are_passed_without_changing_shared_limits(bank_files, monkeypatch):
    state = bank_files
    monkeypatch.setenv("USER_MAX_CHUNKS_PER_DOC", "1")
    monkeypatch.setenv("USER_MAX_TABLE_ROWS", "2")
    get_settings.cache_clear()
    sid = new_session(state)
    upload(state, sid, ("Long sentence. " * 500).encode())
    assert len(state.embedding_batches[-1]) == 1
    assert any("only the first 1" in warning for warning in state.docs.rows[0]["warnings"])
    result = load_file("a.csv", b"item,value\none,1\ntwo,2\nthree,3", max_table_rows=2)
    assert any("first 2" in warning for warning in result.warnings)
    row_text = "\n".join(row.text for row in result.sections if row.kind == "table")
    assert "three" not in row_text and "two" in row_text
    uncapped = load_file("a.csv", b"item,value\none,1\ntwo,2\nthree,3")
    assert not uncapped.warnings


@pytest.mark.parametrize("extension", [".md", ".docx", ".xlsx", ".pptx"])
def test_row_limit_parameter_reaches_table_loaders(bank_files, extension):
    if extension == ".md":
        data = b"| Item | Value |\n|---|---|\n| one | 1 |\n| two | 2 |"
    else:
        buffer = io.BytesIO()
        if extension == ".docx":
            from docx import Document
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            for row, values in zip(table.rows, (("Item","Value"),("one","1"),("two","2"))):
                for cell, value in zip(row.cells, values): cell.text = value
            doc.save(buffer)
        elif extension == ".xlsx":
            from openpyxl import Workbook
            book = Workbook()
            for row in (("Item","Value"),("one",1),("two",2)): book.active.append(row)
            book.save(buffer)
        else:
            from pptx import Presentation
            from pptx.util import Inches
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            table = slide.shapes.add_table(3,2, Inches(1), Inches(1), Inches(5), Inches(2)).table
            for i, row in enumerate((("Item","Value"),("one","1"),("two","2"))):
                for j, value in enumerate(row): table.cell(i,j).text = value
            deck.save(buffer)
        data = buffer.getvalue()
    result = load_file("table" + extension, data, max_table_rows=1)
    assert any("first 1" in warning for warning in result.warnings)


@pytest.mark.parametrize("status", ["failed", "processing"])
def test_failed_or_wait_timeout_never_calls_answer(bank_files, monkeypatch, status):
    state = bank_files
    sid = new_session(state)
    row = upload(state, sid)
    state.docs.rows[0].update(status=status, error="PRIVATE raw provider error")
    monkeypatch.setenv("ATTACHMENT_WAIT_SECONDS", "0")
    get_settings.cache_clear()
    response = state.client.post("/chat", headers=state.headers("a"), json={
        "message":"Read this", "session_id":sid, "attachment_ids":[row["doc_id"]], "stream":False})
    assert "error" in response.json()
    assert "PRIVATE" not in response.text
    assert state.answer_sources == [] and state.messages.rows == []


def test_processing_attachment_emits_status_then_completes(bank_files, monkeypatch):
    state = bank_files
    sid = new_session(state)
    row = upload(state, sid)
    state.docs.rows[0]["status"] = "processing"
    async def finish(delay): state.docs.rows[0]["status"] = "ready"
    monkeypatch.setattr(attachments.asyncio, "sleep", finish)
    response = state.client.post("/chat", headers=state.headers("a"), json={
        "message":"Read this", "session_id":sid, "attachment_ids":[row["doc_id"]], "stream":True})
    assert "event: status" in response.text and "Waiting for attached files" in response.text
    assert "event: done" in response.text and len(state.answer_sources) == 1


def test_delete_during_embedding_does_not_leave_orphan_chunks(bank_files, monkeypatch):
    state = bank_files
    sid = new_session(state)
    row, _ = asyncio.run(document_store.register_document("a.txt", "text/plain", b"fact", scope="user", owner_id="a", session_id=sid))
    async def embed_and_delete(texts):
        await memory.clear_session(sid, user_id="a")
        return [[1.,0.] for _ in texts]
    monkeypatch.setattr(ingestion, "embed_documents", embed_and_delete)
    asyncio.run(document_store.ingest_scoped(str(row["_id"]), "a.txt", b"fact", scope="user", owner_id="a", session_id=sid))
    assert state.docs.rows == [] and state.chunks.rows == []


def test_failed_duplicate_can_be_retried_without_consuming_slot(bank_files):
    state = bank_files
    sid = new_session(state)
    row = upload(state, sid)
    state.docs.rows[0]["status"] = "failed"
    retry = upload(state, sid)
    assert retry["doc_id"] == row["doc_id"] and not retry["duplicate"]
    assert state.sessions.rows[0]["attachment_count"] == 1
    assert state.docs.rows[0]["status"] == "ready"


def test_kb_seed_five_documents_table_and_safe_replacement(bank_files, monkeypatch):
    import sys
    async def noop(): pass
    for name in ("close_mongo", "connect_to_mongo", "ensure_indexes"):
        monkeypatch.setattr(sys.modules["app.db"], name, noop, raising=False)
    from scripts import seed_kb
    state = bank_files
    monkeypatch.setattr(seed_kb, "documents_col", lambda: state.docs)
    monkeypatch.setattr(seed_kb, "chunks_col", lambda: state.chunks)
    sid = new_session(state)
    own = upload(state, sid, b"private", "savings_terms.md")
    asyncio.run(seed_kb.seed_kb())
    first_ids = {row["_id"] for row in state.docs.rows if row.get("scope") == "shared"}
    assert len(first_ids) == 5
    assert any(row["kind"] == "table" and row["scope"] == "shared" for row in state.chunks.rows)
    asyncio.run(seed_kb.seed_kb())
    shared = [row for row in state.docs.rows if row.get("scope") == "shared"]
    assert len(shared) == 5 and all(row["status"] == "ready" for row in shared)
    assert first_ids.isdisjoint({row["_id"] for row in shared})
    assert any(str(row["_id"]) == own["doc_id"] for row in state.docs.rows)
    assert all(row["owner_id"] is None and row["session_id"] is None and row["expires_at"] is None for row in shared)


def test_vector_definition_matches_copyable_atlas_json(bank_files):
    import importlib.util
    root = Path(__file__).parents[1]
    spec = importlib.util.spec_from_file_location("phase3_real_db", root / "app" / "db.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    assert module.vector_index_definition() == json.loads((root / "atlas" / "chunks_vector_index.json").read_text())


def test_global_routes_removed_and_request_scope_is_not_client_controlled(bank_files):
    state = bank_files
    assert state.client.get("/documents", headers=state.headers("a")).status_code == 404
    assert state.client.post("/documents", headers=state.headers("a")).status_code == 404
    sid = new_session(state)
    for extra in ({"doc_ids":[]}, {"owner_id":"b"}, {"scope":"shared"}):
        response = state.client.post("/chat", headers=state.headers("a"), json={"message":"hello", "session_id":sid, **extra})
        assert response.status_code == 422
    assert state.client.post("/chat", headers=state.headers("a"), json={"message":"hello", "session_id":sid,
        "attachment_ids":["invalid"]}).status_code == 404
